"""
app/services/report_service.py
----------------------------------
PDF and Excel report generation for complaints (Advanced Search & Filters
results, or the full complaint set) — the same filtering as GET
/complaints, but rendered as a downloadable report instead of JSON.

PDF: reportlab. Excel: openpyxl. Both are pure-Python, no system deps.
"""
import io
from typing import List, Optional

from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.repository import complaint_repository
from app.utils.timeutil import utcnow

COLUMNS = [
    ("complaint_id", "Ticket ID"),
    ("severity", "Severity"),
    ("num_potholes", "Potholes"),
    ("confidence", "Max Conf."),
    ("status", "Status"),
    ("department", "Department"),
    ("latitude", "Lat"),
    ("longitude", "Lng"),
    ("date_reported", "Filed"),
    ("resolved_at", "Resolved"),
    ("report_count", "Reports"),
]


def _rows(db: Session, severity, status, department, date_from, date_to, search) -> List:
    # limit=10000 — a report is meant to cover "everything matching the
    # filter", not one paginated page of it.
    _, items = complaint_repository.search(
        db, severity=severity, status=status, department=department,
        date_from=date_from, date_to=date_to, search_text=search, limit=10000, offset=0,
    )
    return items


def _cell(c, key):
    val = getattr(c, key)
    if key == "confidence" and val is not None:
        return f"{val * 100:.1f}%"
    if key in ("date_reported", "resolved_at") and val is not None:
        return val.strftime("%Y-%m-%d %H:%M")
    if key in ("latitude", "longitude") and val is not None:
        return f"{val:.5f}"
    return "" if val is None else str(val)


def generate_pdf(db: Session, **filters) -> StreamingResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    items = _rows(db, **filters)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    active_filters = ", ".join(f"{k}={v}" for k, v in filters.items() if v) or "none"
    story = [
        Paragraph("RoadWatch — Complaint Report", styles["Title"]),
        Paragraph(f"Generated {utcnow().strftime('%Y-%m-%d %H:%M UTC')} · Filters: {active_filters} · {len(items)} complaint(s)", styles["Normal"]),
        Spacer(1, 14),
    ]

    header = [label for _, label in COLUMNS]
    data = [header] + [[_cell(c, key) for key, _ in COLUMNS] for c in items]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16181a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)
    doc.build(story)
    buf.seek(0)

    return StreamingResponse(
        buf, media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=roadwatch_report.pdf"},
    )


def generate_excel(db: Session, **filters) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    items = _rows(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints"

    header = [label for _, label in COLUMNS]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16181A")

    for c in items:
        ws.append([_cell(c, key) for key, _ in COLUMNS])

    for i, (key, label) in enumerate(COLUMNS, start=1):
        width = max(len(label) + 2, 12)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    by_sev = {"Low": 0, "Medium": 0, "High": 0, "Critical": 0}
    by_status = {}
    for c in items:
        if c.severity in by_sev:
            by_sev[c.severity] += 1
        by_status[c.status] = by_status.get(c.status, 0) + 1
    ws2.append(["Generated", utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    ws2.append(["Total complaints", len(items)])
    ws2.append([])
    ws2.append(["Severity", "Count"])
    for k, v in by_sev.items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Status", "Count"])
    for k, v in by_status.items():
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=roadwatch_report.xlsx"},
    )
